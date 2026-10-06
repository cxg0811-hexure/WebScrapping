# Generated with GitHub Copilot - [Tracking ID: Hexure-Copilot]
"""Build the formatted Excel workbook (sorted by price, ascending)."""
from __future__ import annotations

import os
import tempfile
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F1F1F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
INR_FORMAT = '"₹"#,##,##0'

COLUMNS = [
    ("#", 6), ("Product", 70), ("Category", 14), ("Store", 16), ("Price (₹)", 14),
    ("MRP (₹)", 14), ("Discount %", 11), ("Sale Price (₹)", 15), ("Sale Discount %", 12),
    ("With Bank Offer (₹)", 18), ("Sale Name", 24), ("Sale Status", 11), ("Sale Starts", 18),
    ("Sale Ends", 18), ("Link", 12), ("Scraped At", 20),
]
SALE_COLUMNS = [
    ("#", 6), ("Product", 70), ("Category", 14), ("Store", 16), ("Sale Price (₹)", 15),
    ("With Bank Offer (₹)", 18), ("Regular Price (₹)", 16), ("MRP (₹)", 14), ("Sale Name", 24),
    ("Sale Status", 11), ("Sale Starts", 18), ("Sale Ends", 18), ("Link", 12),
]
CATEGORY_COLUMNS = [
    ("Product", 70), ("Store", 16), ("Price (₹)", 14), ("MRP (₹)", 14), ("Sale Price (₹)", 15),
    ("With Bank Offer (₹)", 18), ("Sale Name", 24), ("Link", 12),
]


def _money(value):
    return (value, INR_FORMAT) if value else None


def _pct(value):
    return (value / 100, "0.0%") if value is not None else None


def _link(url):
    return ("Open", url)


def _write_row(ws, r, values):
    """Write a row; each value is plain, None, (number, format) or ('Open', url) for hyperlinks."""
    for c, v in enumerate(values, start=1):
        if v is None:
            continue
        if isinstance(v, tuple) and v[0] == "Open":
            cell = ws.cell(r, c, "Open")
            cell.hyperlink, cell.style = v[1], "Hyperlink"
        elif isinstance(v, tuple):
            ws.cell(r, c, v[0]).number_format = v[1]
        else:
            ws.cell(r, c, v)


def _mrp(p):
    return _money(p["mrp"]) if p.get("mrp") and p["mrp"] > p["price"] else None


def _write_header(ws, columns):
    for idx, (title, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"


def build_workbook(products: list[dict], path: Path, generated_at: datetime) -> Path:
    rows = sorted(products, key=lambda p: (p["price"], p["name"]))
    wb = Workbook()

    ws = wb.active
    ws.title = "All Phones"
    _write_header(ws, COLUMNS)
    for i, p in enumerate(rows, start=1):
        _write_row(ws, i + 1, [
            i, p["name"], p["category"], p["store"], _money(p["price"]), _mrp(p), _pct(p.get("discount_pct")),
            _money(p.get("sale_price")), _pct(p.get("sale_discount_pct")), _money(p.get("offer_price")),
            p.get("sale_name", ""), p.get("sale_status", ""), p.get("sale_starts", ""), p.get("sale_ends", ""),
            _link(p["url"]), p["scraped_at"].replace("T", " "),
        ])
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"

    # Sale listings only (e.g. Big Billion Days), sorted by sale price ascending.
    deals = sorted((p for p in rows if p.get("sale_price")), key=lambda p: (p["sale_price"], p["name"]))
    sws = wb.create_sheet("Sale Deals")
    _write_header(sws, SALE_COLUMNS)
    for i, p in enumerate(deals, start=1):
        _write_row(sws, i + 1, [
            i, p["name"], p["category"], p["store"], _money(p["sale_price"]), _money(p.get("offer_price")),
            _money(p["price"]), _mrp(p), p.get("sale_name", ""), p.get("sale_status", ""),
            p.get("sale_starts", ""), p.get("sale_ends", ""), _link(p["url"]),
        ])
    if deals:
        sws.auto_filter.ref = f"A1:{get_column_letter(len(SALE_COLUMNS))}{len(deals) + 1}"

    # One sheet per category, also sorted ascending by price.
    by_cat: dict[str, list[dict]] = defaultdict(list)
    for p in rows:
        by_cat[p["category"]].append(p)
    for cat in sorted(by_cat):
        cws = wb.create_sheet(cat[:31])
        _write_header(cws, CATEGORY_COLUMNS)
        for r, p in enumerate(by_cat[cat], start=2):
            _write_row(cws, r, [
                p["name"], p["store"], _money(p["price"]), _mrp(p), _money(p.get("sale_price")),
                _money(p.get("offer_price")), p.get("sale_name", ""), _link(p["url"]),
            ])

    summary = wb.create_sheet("Summary", 0)
    summary["A1"] = "Phone Price Report (iPhone, Samsung, Google Pixel)"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A2"] = f"Generated: {generated_at:%Y-%m-%d %H:%M:%S}"
    summary["A3"] = f"Total listings: {len(rows)}  |  Sale listings: {len(deals)} (see 'Sale Deals' sheet)"
    _hdr = [("Category", 18), ("Listings", 10), ("Lowest Price (₹)", 18), ("Cheapest Product", 70), ("Store", 16)]
    for idx, (title, width) in enumerate(_hdr, start=1):
        cell = summary.cell(5, idx, title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        summary.column_dimensions[get_column_letter(idx)].width = width
    for r, cat in enumerate(sorted(by_cat, key=lambda c: by_cat[c][0]["price"]), start=6):
        cheapest = by_cat[cat][0]
        summary.cell(r, 1, cat)
        summary.cell(r, 2, len(by_cat[cat]))
        summary.cell(r, 3, cheapest["price"]).number_format = INR_FORMAT
        summary.cell(r, 4, cheapest["name"])
        summary.cell(r, 5, cheapest["store"])

    # Write atomically so a download never receives a half-written file.
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=path.parent)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return path
