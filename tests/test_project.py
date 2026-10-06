# Generated with GitHub Copilot - [Tracking ID: Hexure-Copilot]
import json
from datetime import datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import config
from core.excel import build_workbook
from core.filters import categorize, is_target_phone, normalize_name
from scrapers.base import Product, parse_price
from scrapers.flipkart import FlipkartScraper
from scrapers.poorvika import PoorvikaScraper
from scrapers.reliance_digital import RelianceDigitalScraper
from scrapers.vijay_sales import VijaySalesScraper


class FakeResponse:
    status_code = 200

    def __init__(self, payload=None, text=""):
        self._payload, self.text = payload, text

    def json(self):
        return self._payload


@pytest.mark.parametrize("text,expected", [
    ("₹1,43,990", 143990.0), ("₹ 25,499.50", 25499.5), ("free", None), (None, None), ("₹0", None),
])
def test_parse_price(text, expected):
    assert parse_price(text) == expected


@pytest.mark.parametrize("title,ok", [
    ("Apple iPhone 17 256GB", True),
    ("Samsung Galaxy S25 Ultra 5G (256GB)", True),
    ("Google Pixel 9 Pro (128GB)", True),
    ("Samsung Galaxy Watch 7", False),
    ("Google Pixel Buds Pro", False),
    ("Samsung Galaxy Tab S9", False),
    ("Apple iPhone 16 Silicone Case with MagSafe", False),
    ("Spigen Tempered Glass for iPhone 16", False),
    ("Apple iPhone 15 Refurbished", False),
])
def test_genuine_filter(title, ok):
    assert is_target_phone(title) is ok


def test_categorize_and_normalize():
    assert categorize("Samsung Galaxy S25", "x") == "Samsung"
    assert categorize("Google Pixel 9", "x") == "Google Pixel"
    assert categorize("Apple iPhone 17", "x") == "iPhone"
    assert normalize_name("Apple iPhone Air\u200b: 13-inch with M3 chip") == "Apple iPhone Air"


def test_reliance_parser(monkeypatch):
    payload = {"items": [
        {"name": "Apple iPhone 18 Pro", "slug": "apple-18", "uid": 1, "sellable": True,
         "brand": {"name": "Apple"}, "price": {"effective": {"min": 150000}, "marked": {"min": 160000}}},
        {"name": "OnePlus Phone", "slug": "s", "brand": {"name": "OnePlus"}, "price": {"effective": {"min": 1}}},
    ]}
    s = RelianceDigitalScraper()
    monkeypatch.setattr(s, "fetch", lambda *a, **k: FakeResponse(payload))
    items = list(s.search("apple iphone", "iPhone", 1))
    assert len(items) == 1 and items[0].price == 150000 and items[0].discount_pct == 6.2


def test_vijay_parser(monkeypatch):
    payload = {"response": {"products": [
        {"brand": ["Apple"], "title": "Apple iPhone 17", "productUrl": "https://x/p/1", "uniqueId": "1",
         "cityId_1_sellingPrice_unx_d": 80000, "mrp": 82000, "cityId_1_status_unx_ts": "Available"},
        {"brand": ["Apple"], "title": "Apple iPhone 16", "productUrl": "https://x/p/2", "uniqueId": "2",
         "price": 70000, "cityId_1_status_unx_ts": "Out of Stock"},
    ]}}
    s = VijaySalesScraper()
    monkeypatch.setattr(s, "fetch", lambda *a, **k: FakeResponse(payload))
    items = list(s.search("apple iphone", "iPhone", 1))
    assert [p.price for p in items] == [80000.0]


def test_poorvika_parser(monkeypatch):
    doc = {"name": "Apple iPhone 17 ( Black, 256GB )", "slug": "apple-iphone-17", "id": "a",
           "special_price": 99900, "mrp": [{"price": [{"price": 109900}]}]}
    payload = {"results": [{"hits": [{"document": doc}]}]}
    s = PoorvikaScraper()
    monkeypatch.setattr(s.session, "post", lambda *a, **k: FakeResponse(payload))
    item = next(iter(s.search("apple iphone", "iPhone", 1)))
    assert item.url.endswith("/apple-iphone-17/p") and item.mrp == 109900.0


def test_flipkart_skips_exchange_offer(monkeypatch):
    html = """<div data-id="X1"><a href="/apple-iphone-16/p/itm?pid=X1" title="Apple iPhone 16 (Black, 128 GB)">
    <div><div>₹69,900</div><div><div>₹ 79,900</div></div></div>
    <div>Upto <span>₹47,000</span> Off on Exchange</div></a></div>"""
    s = FlipkartScraper()
    monkeypatch.setattr(s, "fetch", lambda *a, **k: FakeResponse(text=html))
    item = next(iter(s.search("apple iphone", "iPhone", 1)))
    assert item.price == 69900 and item.mrp == 79900


def test_scrape_dedupes_and_filters(monkeypatch):
    s = RelianceDigitalScraper()
    mk = lambda n, p: Product(name=n, price=p, store="S", category="iPhone", url="u", product_id="id1")
    monkeypatch.setattr(s, "search", lambda q, c, pg: [mk("Apple iPhone 17", 90), mk("Apple iPhone 17", 80),
                                                       Product("Apple iPhone Case", 5, "S", "x", "u2", product_id="c")] if pg == 1 else [])
    monkeypatch.setattr(s, "polite_pause", lambda: None)
    out = s.scrape({"iPhone": "apple iphone"}, 2)
    assert len(out) == 1 and out[0].price == 80


def test_excel_sorted_ascending(tmp_path):
    rows = [
        Product("Apple iPhone 17", 90000, "A", "iPhone", "http://a", mrp=95000).to_dict(),
        Product("Samsung Galaxy S25", 12000, "B", "Samsung", "http://b", sale_price=10000, sale_name="Sale").to_dict(),
        Product("Google Pixel 9", 40000, "C", "Google Pixel", "http://c").to_dict(),
    ]
    path = build_workbook(rows, tmp_path / "t.xlsx", datetime.now())
    ws = load_workbook(path)["All Phones"]
    prices = [ws.cell(r, 5).value for r in range(2, ws.max_row + 1)]
    assert prices == [12000, 40000, 90000]


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "LATEST_EXCEL", tmp_path / "latest.xlsx")
    monkeypatch.setattr(config, "LATEST_JSON", tmp_path / "latest.json")
    import app as app_module
    return TestClient(app_module.app)


def test_download_404_when_no_report(client):
    assert client.get("/download").status_code == 404


def test_download_and_api(client):
    rows = [Product("Apple iPhone 17", 90000, "A", "iPhone", "http://a").to_dict()]
    build_workbook(rows, config.LATEST_EXCEL, datetime.now())
    config.LATEST_JSON.write_text(json.dumps({"generated_at": "2026-01-01T00:00:00", "stores": {}, "products": rows}))
    r = client.get("/download")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert client.get("/api/products", params={"q": "iphone"}).json()["count"] == 1
    assert client.get("/api/products", params={"store": "zzz"}).json()["count"] == 0
    assert client.get("/api/status").json()["excel_available"] is True


def test_sale_discount():
    p = Product("Samsung Galaxy S25", 80000, "S", "Samsung", "u", sale_price=60000, sale_name="BBD")
    assert p.sale_discount_pct == 25.0
    assert Product("x", 5, "S", "c", "u").sale_discount_pct is None


def test_flipkart_sale_tag(monkeypatch):
    html = """<div data-id="X2"><a href="/s/p/itm?pid=X2" title="Samsung Galaxy A36 5G (128 GB)">x</a>
    <div>Coming Soon</div><div>₹30,999</div><div>₹38,999</div><div>Big Billion Days Price</div></div>"""
    s = FlipkartScraper()
    monkeypatch.setattr(s, "fetch", lambda *a, **k: FakeResponse(text=html))
    item = next(iter(s.search("samsung galaxy", "Samsung", 1)))
    assert item.sale_price == 30999.0 and item.sale_status == "Upcoming" and "Big Billion" in item.sale_name
